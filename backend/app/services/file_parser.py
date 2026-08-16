import io
import os


def extract_text_from_resume_file(filename: str, content: bytes) -> str:
    ext = os.path.splitext(filename)[1].lower()

    if ext == ".txt":
        text = content.decode("utf-8", errors="replace")

    elif ext == ".pdf":
        try:
            import pypdf
        except ImportError:
            raise ValueError("pypdf がインストールされていません。pip install pypdf を実行してください")
        try:
            reader = pypdf.PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception as e:
            raise ValueError(f"PDFの読み込みに失敗しました: {e}")

    elif ext == ".docx":
        try:
            import docx
        except ImportError:
            raise ValueError("python-docx がインストールされていません。pip install python-docx を実行してください")
        try:
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception as e:
            raise ValueError(f"Wordファイルの読み込みに失敗しました: {e}")

    elif ext in (".xlsx", ".xls"):
        try:
            import openpyxl
        except ImportError:
            raise ValueError("openpyxl がインストールされていません")
        try:
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = " ".join(str(c) for c in row if c is not None)
                    if line.strip():
                        lines.append(line)
            text = "\n".join(lines)
        except Exception as e:
            raise ValueError(f"Excelファイルの読み込みに失敗しました: {e}")

    else:
        raise ValueError(f"未対応の形式です: {ext}（対応形式: .txt / .pdf / .xlsx / .docx）")

    if not text.strip():
        raise ValueError(
            "テキストを抽出できませんでした。"
            "スキャン画像PDFの場合は抽出不可です。テキスト形式のPDFをご使用ください。"
        )

    return text.strip()
